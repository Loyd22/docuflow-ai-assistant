"""
Tests for Microsoft Word DOCX text extraction.
"""

from pathlib import Path

import pytest
from docx import Document

from app.processing.docx_extractor import DOCXExtractor


def test_extract_returns_docx_paragraph_text(tmp_path: Path) -> None:
    """The extractor should return text from DOCX paragraphs."""

    file_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("DocuFlow Document Processing")
    document.add_paragraph("This is a DOCX extraction test.")
    document.save(file_path)

    extractor = DOCXExtractor()

    result = extractor.extract(file_path)

    assert result == ("DocuFlow Document Processing\nThis is a DOCX extraction test.")


def test_extract_ignores_empty_paragraphs(tmp_path: Path) -> None:
    """The extractor should ignore empty DOCX paragraphs."""

    file_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("")
    document.add_paragraph("Second paragraph")
    document.save(file_path)

    extractor = DOCXExtractor()

    result = extractor.extract(file_path)

    assert result == "First paragraph\nSecond paragraph"


def test_extract_raises_error_when_file_does_not_exist(
    tmp_path: Path,
) -> None:
    """The extractor should reject missing DOCX files."""

    file_path = tmp_path / "missing.docx"

    extractor = DOCXExtractor()

    with pytest.raises(FileNotFoundError):
        extractor.extract(file_path)


def test_extract_raises_error_when_path_is_directory(
    tmp_path: Path,
) -> None:
    """The extractor should reject directory paths."""

    extractor = DOCXExtractor()

    with pytest.raises(ValueError):
        extractor.extract(tmp_path)
